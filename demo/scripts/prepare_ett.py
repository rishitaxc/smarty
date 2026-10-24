"""Real-data case: two power transformers from the public ETT dataset (Zhou et al., "Informer", AAAI 2021).

Downloads ETTh1.csv and ETTh2.csv (hourly, Jul 2016 - Jun 2018) from github.com/zhouhaoyi/ETDataset,
averages them per day for one year (1 Jul 2016 - 30 Jun 2017) and writes:
  demo/ETT_transformers_daily.xlsx   (the file the agents open)
  demo/public/data/sheet_ett.json    (sheet.html?ds=ett)
Nothing is injected: the event Scout should find (TX-2 overload, 18 Apr 2017) is in the published data.
Run: python3 -I demo/scripts/prepare_ett.py
"""
import csv, io, json, os, ssl, statistics, urllib.request
from collections import defaultdict
from datetime import date, timedelta
from openpyxl import Workbook
from openpyxl.styles import Font

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
OUT = os.path.join(ROOT, "public", "data")
BASE = "https://raw.githubusercontent.com/zhouhaoyi/ETDataset/main/ETT-small/"
TLS = ssl.create_default_context(cafile="/etc/ssl/cert.pem") if os.path.exists("/etc/ssl/cert.pem") else None
START, DAYS = date(2016, 7, 1), 365
UNITS = {"TX-1": "ETTh1.csv", "TX-2": "ETTh2.csv"}
FILE = "ETT_transformers_daily.xlsx"


def daily(name):
    text = urllib.request.urlopen(BASE + name, timeout=60, context=TLS).read().decode()
    by_day = defaultdict(list)
    for r in csv.DictReader(io.StringIO(text)):
        by_day[r["date"][:10]].append(r)
    return by_day


def main():
    days = [START + timedelta(days=i) for i in range(DAYS)]
    tabs = [("Oil temp (°C)", "oil", lambda rs: statistics.mean(float(r["OT"]) for r in rs), 1),
            ("Peak oil temp (°C)", "peak", lambda rs: max(float(r["OT"]) for r in rs), 1),
            ("Load (HUFL)", "load", lambda rs: statistics.mean(float(r["HUFL"]) for r in rs), 1)]
    out = {"labels": [d.strftime("%d %b") for d in days], "turbines": list(UNITS), "file": FILE, "domain": "Energy", "tabs": [],
           "source": "Real data: ETT dataset (Zhou et al., Informer, AAAI 2021) · two power transformers, China · "
                     "hourly readings averaged per day, 1 Jul 2016 – 30 Jun 2017 · github.com/zhouhaoyi/ETDataset"}
    raw = {u: daily(f) for u, f in UNITS.items()}
    wb = Workbook(); wb.remove(wb.active)
    for title, k, agg, dec in tabs:
        out["tabs"].append({"k": k, "title": title})
        out[k] = {u: [round(agg(raw[u][d.isoformat()]), dec) for d in days] for u in UNITS}
        ws = wb.create_sheet(title.split(" (")[0][:31])
        ws.append(["Date"] + list(UNITS))
        for i, lb in enumerate(out["labels"]):
            ws.append([lb] + [out[k][u][i] for u in UNITS])
        for row in ws.iter_rows():
            for c in row:
                c.font = Font(name="Arial", size=10, bold=c.row == 1)
    src = wb.create_sheet("Source")
    for line in [["Dataset", "ETT (Electricity Transformer Temperature), ETTh1 + ETTh2"],
                 ["Paper", "Zhou et al., Informer, AAAI 2021"], ["URL", "https://github.com/zhouhaoyi/ETDataset"],
                 ["Columns", "OT = oil temperature (°C); HUFL = high useful load, as published"],
                 ["Processing", "Hourly values averaged per day (peak = daily max). No values changed or added."]]:
        src.append(line)
    wb.save(os.path.join(ROOT, FILE))
    json.dump(out, open(os.path.join(OUT, "sheet_ett.json"), "w"), separators=(",", ":"))
    i = [d.isoformat() for d in days].index("2017-04-18")
    print(f"wrote {FILE}: {DAYS} days × {len(UNITS)} transformers · 18 Apr 2017 TX-2 oil {out['oil']['TX-2'][i]} °C, load {out['load']['TX-2'][i]}")


if __name__ == "__main__":
    main()
